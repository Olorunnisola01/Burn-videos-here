from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

root=Path('Relationships psychology/The Human Pattern'); src=root/'09_COMPLETE_MANUSCRIPT.md'; assets=root/'assets'; out=root/'The Human Pattern - Amazon Ready Interior.docx'
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.8); sec.left_margin=Inches(.9); sec.right_margin=Inches(.75)
styles=doc.styles
styles['Normal'].font.name='Georgia'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor(35,48,62); styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.12
for name,size,color in [('Title',34,'17324D'),('Subtitle',16,'2B6777'),('Heading 1',23,'17324D'),('Heading 2',16,'2B6777'),('Heading 3',12,'B05B3B'),('Caption',9,'64748B')]:
 s=styles[name]; s.font.name='Aptos Display' if name!='Caption' else 'Aptos'; s.font.size=Pt(size); s.font.bold=name!='Caption'; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(14); s.paragraph_format.space_after=Pt(8)

def field(paragraph, instruction):
 run=paragraph.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),instruction); run._r.addnext(fld)
def page(): doc.add_page_break()
def caption(text):
 p=doc.add_paragraph(style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(text)
def insert_fig(num, caption_text):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(assets / f"figure_{num}_1_{ {'1':'behavior_map','2':'behavior_loop','3':'emotion_decision','4':'perception','5':'worth_sources','6':'attachment_cycle','7':'observation_framework'}[str(num)] }.png"), width=Inches(6.3))
 caption(f'Figure {num}.1. {caption_text}')
figcaps={1:'Multiple psychological pathways can produce the same observable behavior.',2:'Immediate rewards can reinforce behavioral patterns that create later costs.',3:'Emotional information and practical evidence can be integrated into a values-aligned decision.',4:'Perception is shaped by attention, expectation, memory, culture, and interpretation.',5:'A resilient sense of dignity draws from more than one source of worth.',6:'Attachment cycles can be interrupted through clear needs, planned space, and reconnection.',7:'Better observation begins with the event and examines the forces around it.'}
lines=src.read_text().splitlines(); skip=0; chapter=None; first=True
for idx,line in enumerate(lines):
 s=line.strip().replace('**','').replace('`','')
 if skip: skip-=1; continue
 if not s: continue
 if s=='---': page(); continue
 if s.startswith('[GRAPHICAL ELEMENT]'):
  # skip following description line; actual figure is inserted after the chapter content below
  skip=1; continue
 if s.startswith('# Part '):
  page(); p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(s[2:]); r.bold=True; r.font.size=Pt(26); r.font.color.rgb=RGBColor(23,50,77); continue
 if s.startswith('# Chapter '):
  if chapter in figcaps and chapter<=7: insert_fig(chapter, figcaps[chapter])
  m=re.search(r'Chapter (\d+)',s); chapter=int(m.group(1)) if m else None; page(); doc.add_paragraph(s[2:],style='Heading 1'); continue
 if s.startswith('# '):
  text=s[2:]
  if first:
   p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(text); r.bold=True; r.font.size=Pt(34); r.font.color.rgb=RGBColor(23,50,77); first=False; page()
  else: doc.add_paragraph(text,style='Heading 1')
 elif s.startswith('## '): doc.add_paragraph(s[3:],style='Heading 2')
 elif s.startswith('### '): doc.add_paragraph(s[4:],style='Heading 3')
 elif re.match(r'^\d+\. ',s): doc.add_paragraph(re.sub(r'^\d+\. ','',s),'List Number')
 elif s.startswith('- '): doc.add_paragraph(s[2:],'List Bullet')
 elif s.startswith('**') and s.endswith('**'):
  p=doc.add_paragraph(); r=p.add_run(s.strip('*')); r.bold=True
 elif s.startswith('|'):
  cells=[x.strip() for x in s.strip('|').split('|')]
  if all(set(x)<=set('-: ') for x in cells): continue
  t=doc.add_table(rows=1,cols=len(cells)); t.style='Light Shading Accent 1'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
  for c,v in zip(t.rows[0].cells,cells): c.text=v
 else:
  p=doc.add_paragraph(); parts=re.split(r'(\*\*.*?\*\*)',s)
  for part in parts:
   r=p.add_run(part[2:-2] if part.startswith('**') and part.endswith('**') else part); r.bold=part.startswith('**') and part.endswith('**')
if chapter in figcaps and chapter<=7: insert_fig(chapter, figcaps[chapter])
# real TOC field at first occurrence is not easy after parsing; append navigation page with updateable field
# Add a front navigation page at end as a Word-updatable TOC field, plus note for Word update.
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('NAVIGATION').bold=True
p=doc.add_paragraph(); field(p,'TOC \\o "1-3" \\h \\z \\u')
# headers and footers
for section in doc.sections:
 header=section.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; rr=header.add_run('THE HUMAN PATTERN'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(100,116,139)
 footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=footer.add_run('Adele Olorunnisola  •  '); rr.font.size=Pt(8); field(footer,'PAGE')
doc.core_properties.title='The Human Pattern — Amazon-Ready Interior'; doc.core_properties.author='Adele Olorunnisola'
doc.save(out); print(out, out.stat().st_size)
