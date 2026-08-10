from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

src=Path('Relationships psychology/The Human Pattern/09_COMPLETE_MANUSCRIPT.md')
out=Path('Relationships psychology/The Human Pattern/The Human Pattern - Complete Manuscript.docx')
doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.8); sec.left_margin=Inches(.9); sec.right_margin=Inches(.9)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor(35,48,62)
styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.13
for name,size,color in [('Title',34,'17324D'),('Heading 1',23,'17324D'),('Heading 2',16,'2B6777'),('Heading 3',12,'B05B3B')]:
 s=styles[name]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
 s.paragraph_format.space_before=Pt(14); s.paragraph_format.space_after=Pt(8)

def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def page_break(): doc.add_page_break()
def add_para(text,style=None):
 p=doc.add_paragraph(style=style); p.add_run(text); return p

def add_callout(text):
 t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
 c=t.cell(0,0); shade(c,'EAF2F4'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(text); r.bold=True; r.font.color.rgb=RGBColor(43,103,119)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=src.read_text().splitlines(); in_list=False; first_title=True
for line in lines:
 s=line.strip()
 if not s: continue
 if s=='---': page_break(); continue
 if s.startswith('[GRAPHICAL ELEMENT]'):
  add_callout('GRAPHICAL ELEMENT\n'+(lines[lines.index(line)+1].strip() if False else 'Text-based diagram or callout for the typeset edition.'))
  continue
 if s.startswith('# '):
  text=s[2:]
  if first_title:
   p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(text); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(34); r.font.color.rgb=RGBColor(23,50,77)
   first_title=False; page_break()
  else: doc.add_paragraph(text,style='Heading 1')
 elif s.startswith('## '): doc.add_paragraph(s[3:],style='Heading 2')
 elif s.startswith('### '): doc.add_paragraph(s[4:],style='Heading 3')
 elif re.match(r'^\d+\. ',s):
  add_para(re.sub(r'^\d+\. ','',s),'List Number')
 elif s.startswith('- '): add_para(s[2:],'List Bullet')
 elif s.startswith('**') and s.endswith('**'):
  p=doc.add_paragraph(); r=p.add_run(s.strip('*')); r.bold=True
 elif s.startswith('|'):
  # simple markdown table rows
  cells=[x.strip() for x in s.strip('|').split('|')]
  if all(set(x)<=set('-: ') for x in cells): continue
  t=doc.add_table(rows=1, cols=len(cells)); t.style='Light Shading Accent 1'
  for c,v in zip(t.rows[0].cells,cells): c.text=v
 else:
  # light inline emphasis
  p=doc.add_paragraph()
  parts=re.split(r'(\*\*.*?\*\*)',s)
  for part in parts:
   if part.startswith('**') and part.endswith('**'):
    r=p.add_run(part[2:-2]); r.bold=True
   else: p.add_run(part)
# footer
for section in doc.sections:
 footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=footer.add_run('THE HUMAN PATTERN  •  ADELE OLORUNNISOLA'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(120,130,140)
doc.core_properties.title='The Human Pattern'; doc.core_properties.author='Adele Olorunnisola'; doc.core_properties.subject='Psychology of human behavior'
doc.save(out); print(out, out.stat().st_size)
